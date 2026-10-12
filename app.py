import os
import io
import time
import json
import base64
import tempfile
import datetime
import urllib.parse
import streamlit as st
import streamlit.components.v1 as components
from PIL import Image
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from google import genai
from google.genai import types

# Page setup
st.set_page_config(
    page_title="MLC తాతా మధు - AI ప్రెస్ నోట్ & సందేశాల కన్సోల్",
    page_icon="📰",
    layout="wide",
    initial_sidebar_state="collapsed"
)

# Helper to read and encode local files to base64
def get_base64_image(image_path):
    if os.path.exists(image_path):
        with open(image_path, "rb") as img_file:
            return base64.b64encode(img_file.read()).decode()
    return None

# Crop top header from letterhead.png (Top ~18-19%)
def get_cropped_letterhead_banner(image_path):
    if os.path.exists(image_path):
        try:
            with Image.open(image_path) as img:
                width, height = img.size
                header_box = (0, 0, width, int(height * 0.19))
                cropped = img.crop(header_box)
                buffer = io.BytesIO()
                cropped.save(buffer, format="PNG")
                return base64.b64encode(buffer.getvalue()).decode()
        except Exception:
            return get_base64_image(image_path)
    return None

# JavaScript 1-Click Clipboard Copy Component
def render_copy_button(text_to_copy, button_label="📋 క్లిక్‌తో కాపీ చేయండి (Copy to Clipboard)", element_id="copy_btn"):
    safe_text = json.dumps(text_to_copy)
    html_code = f"""
    <div style="margin: 8px 0;">
        <button id="{element_id}" onclick="copyText_{element_id}()" 
            style="background: linear-gradient(90deg, #4c1d95 0%, #7e22ce 100%);
                   color: white; border: none; padding: 10px 18px; border-radius: 8px;
                   font-size: 14px; font-weight: bold; cursor: pointer; display: inline-flex; align-items: center; gap: 8px;">
            {button_label}
        </button>
        <span id="msg_{element_id}" style="color: #15803d; font-weight: bold; margin-left: 12px; display: none;">✅ క్లిప్‌బోర్డ్‌కి కాపీ అయ్యింది!</span>
    </div>
    <script>
    function copyText_{element_id}() {{
        const text = {safe_text};
        navigator.clipboard.writeText(text).then(() => {{
            const msg = document.getElementById("msg_{element_id}");
            msg.style.display = "inline";
            setTimeout(() => {{ msg.style.display = "none"; }}, 2500);
        }}).catch(err => {{
            console.error('Could not copy text: ', err);
        }});
    }}
    </script>
    """
    components.html(html_code, height=52)

bg_image_base64 = get_base64_image("background.png")
dev_image_base64 = get_base64_image("Sumanth.jpg") or get_base64_image("sumanth.jpg")

lh_filename = None
for fname in ["letterhead.png", "letter head(2).jpg", "letterhead.jpg"]:
    if os.path.exists(fname):
        lh_filename = fname
        break

lh_banner_base64 = get_cropped_letterhead_banner(lh_filename) if lh_filename else None

if bg_image_base64:
    bg_style = f"""
    .stApp {{
        background: linear-gradient(135deg, rgba(248, 246, 252, 0.96) 0%, rgba(255, 255, 255, 0.98) 100%),
                    url("data:image/png;base64,{bg_image_base64}");
        background-size: cover;
        background-position: center top;
        background-repeat: no-repeat;
        background-attachment: fixed;
    }}
    """
else:
    bg_style = """
    .stApp {
        background: linear-gradient(135deg, #f8f6fc 0%, #ffffff 50%, #f7f0fc 100%);
        background-attachment: fixed;
    }
    """

st.markdown(f"""
<style>
    @import url('https://fonts.googleapis.com/css2?family=Anek+Telugu:wght@300;400;500;600;700;800&display=swap');
    
    {bg_style}
    
    html, body, p, div:not([data-testid="stIconMaterial"]), h1, h2, h3, h4, h5, h6, input, textarea, button {{
        font-family: 'Anek Telugu', -apple-system, BlinkMacSystemFont, sans-serif !important;
    }}

    span[data-testid="stIconMaterial"], .material-symbols-rounded, .material-symbols-outlined {{
        font-family: 'Material Symbols Rounded', 'Material Symbols Outlined' !important;
        font-size: 22px !important;
        line-height: 1 !important;
        white-space: nowrap !important;
    }}

    section[data-testid="stSidebar"] {{
        background-color: #f7effe !important;
        border-right: 1.5px solid #d8b4fe !important;
        box-shadow: 2px 0 14px rgba(88, 28, 135, 0.05);
    }}

    .header-card {{
        background: linear-gradient(90deg, #3b0764 0%, #581c87 65%, #7e22ce 100%) !important;
        color: #ffffff !important;
        padding: 22px 30px;
        border-radius: 14px;
        box-shadow: 0 8px 24px rgba(59, 7, 100, 0.25);
        margin-bottom: 20px;
        border: 1px solid rgba(255, 255, 255, 0.2);
    }}
    .header-card h1 {{
        color: #ffffff !important;
        font-size: 26px !important;
        font-weight: 800 !important;
        margin: 0 !important;
        line-height: 1.3 !important;
    }}
    .header-card p {{
        color: #f3e8ff !important;
        font-size: 14px !important;
        margin: 6px 0 0 0 !important;
        font-weight: 500 !important;
    }}

    .stTextInput>div>div>input, .stTextArea>div>div>textarea {{
        background-color: #ffffff !important;
        color: #0f172a !important;
        border: 1.5px solid #cbd5e1 !important;
        border-radius: 10px !important;
        font-size: 16px !important;
    }}
    .stTextInput>div>div>input:focus, .stTextArea>div>div>textarea:focus {{
        border-color: #7e22ce !important;
        box-shadow: 0 0 0 2px rgba(126, 34, 206, 0.18) !important;
    }}

    div.stButton > button[kind="primary"] {{
        background: linear-gradient(90deg, #4c1d95 0%, #7e22ce 100%) !important;
        color: #ffffff !important;
        font-size: 16px !important;
        font-weight: 700 !important;
        border: none !important;
        border-radius: 10px !important;
        padding: 12px 24px !important;
        box-shadow: 0 4px 14px rgba(126, 34, 206, 0.35) !important;
        transition: all 0.2s ease !important;
    }}

    .letterhead-container {{
        background-color: #ffffff;
        border: 1.5px solid #e2e8f0;
        border-radius: 8px;
        padding: 35px 48px;
        box-shadow: 0 8px 30px rgba(0,0,0,0.08);
        color: #111111;
        line-height: 1.95;
        font-size: 18px;
        max-width: 860px;
        margin: 0 auto;
        box-sizing: border-box;
    }}

    .letterhead-banner-img {{
        width: 100%;
        max-height: 180px;
        object-fit: contain;
        display: block;
        margin: 0 auto 16px auto;
        border-bottom: 2px solid #581c87;
        padding-bottom: 8px;
    }}

    .dev-badge {{
        display: flex;
        align-items: center;
        gap: 12px;
        background: #ffffff;
        border: 1.5px solid #e9d5ff;
        border-radius: 14px;
        padding: 10px 14px;
        box-shadow: 0 2px 10px rgba(0,0,0,0.04);
        margin-top: 15px;
    }}
    .dev-img {{
        width: 52px;
        height: 52px;
        border-radius: 50%;
        object-fit: cover;
        border: 2px solid #7e22ce;
        flex-shrink: 0;
    }}
    .dev-text {{
        font-size: 13px;
        font-weight: 700;
        color: #1e293b;
        line-height: 1.25;
    }}
    .dev-sub {{
        font-size: 11px;
        color: #6b21a8;
        font-weight: 600;
        margin-top: 3px;
        line-height: 1.3;
    }}
</style>
""", unsafe_allow_html=True)

# ----------------- SYSTEM INSTRUCTIONS -----------------
SYSTEM_INSTRUCTION_BASE = (
    "You are the EXCLUSIVE Chief Media Secretary & Political Communication Director "
    "for Sri Tata Madhusudhan (Tata Madhu) Garu, Member of Legislative Council (MLC), "
    "Bharat Rashtra Samithi (BRS).\n\n"
    "STRICT CONSTRAINTS:\n"
    "1. LEADER EXCLUSIVITY: Every statement, message, meeting invite, critique, or declaration must be strictly attributed "
    "to MLC Tata Madhusudhan (శాసనమండలి సభ్యులు తాతా మధుసూదన్ / తాతా మధు). Under no circumstances generate messages for anyone else.\n"
    "2. SITUATIONAL COMPREHENSION: Whether the user input is in English, Telugu, Tanglish, or bullet points, synthesize it into "
    "authentic, authoritative, standard journalistic & political Telugu.\n"
    "3. JOURNALISTIC & OFFICIAL TELUGU: Write in dignified, authentic, culturally grounded Telugu.\n"
)

SYSTEM_INSTRUCTION_LONG = SYSTEM_INSTRUCTION_BASE + (
    "\nFORMAT: COMPREHENSIVE OFFICIAL PRESS NOTE (సమగ్ర అధికారిక పత్రికా ప్రకటన)\n"
    "- Header: అధికారిక పత్రికా ప్రకటన\n"
    "- స్థలం & తేదీ (Dateline format)\n"
    "- ప్రధాన శీర్షిక (Impactful bold headline highlighting 'ఎమ్మెల్సీ తాతా మధు')\n"
    "- లీడ్ పేరా (Complete background, venue, core political punch and primary statement)\n"
    "- సమగ్ర వివరాలు (2 rich descriptive paragraphs with facts, political context, and arguments)\n"
    "- ముఖ్యాంశాలు (4 to 6 strong bullet points summarizing key critiques or demands)\n"
    "- ముగింపు (Closing political resolution, warning to opposition, or sincere tribute)\n"
    "- విడుదల: ఎమ్మెల్సీ తాతా మధుసూదన్ గారి కార్యాలయం\n"
)

SYSTEM_INSTRUCTION_SHORT = SYSTEM_INSTRUCTION_BASE + (
    "\nFORMAT: CRISP / SHORT PRESS RELEASE (సంక్షిప్త పత్రికా ప్రకటన / డిజిటల్ బులెటిన్)\n"
    "- Header: సంక్షిప్త పత్రికా ప్రకటన\n"
    "- స్థలం & తేదీ\n"
    "- ప్రధాన శీర్షిక (Crisp, punchy headline featuring 'ఎమ్మెల్సీ తాతా మధు')\n"
    "- ప్రధాన ప్రకటన (1-2 sharp paragraphs covering the exact news without filler words)\n"
    "- ముఖ్యాంశాలు (3 crisp, high-impact bullet points)\n"
    "- ముగింపు ప్రకటన (1 final crisp closing sentence)\n"
    "- విడుదల: ఎమ్మెల్సీ తాతా మధుసూదన్ గారి కార్యాలయం\n"
)

SYSTEM_INSTRUCTION_MESSAGES = SYSTEM_INSTRUCTION_BASE + (
    "\nTASK: GENERATE 3 DEDICATED SITUATIONAL & MEETING MESSAGES BASED ON THE INPUT PROMPT:\n"
    "Return the response formatted EXACTLY with these 3 sections separated by specific tags:\n\n"
    "===LEADERS_MSG===\n"
    "(Formal, dignified, authoritative invite/message addressed to BRS District Leaders, Public Representatives, ZPTCs, MPTCs, and senior party functionaries. Mention date, time, venue, agenda importance, and request mandatory attendance).\n\n"
    "===CADRE_MSG===\n"
    "(Inspiring, high-energy, motivational mobilization message for BRS party workers, youth wings, social media warriors, and followers. High energy tone, invoking Telangana sentiments, unity, and party solidarity with slogans like 'జై తెలంగాణ - జై కేసీఆర్ - జై తాతా మధు').\n\n"
    "===PRESS_INVITE===\n"
    "(Official Media Advisory / Press Invitation for Print, Electronic & Digital Media correspondents. Formal invitation header, date, time, venue, topic, request for coverage, and coordinator contact footer).\n"
)

# Transliteration Helper
def google_transliterate_telugu(text: str) -> str:
    if not text.strip():
        return ""
    words = text.split()
    converted_words = []
    url = "https://inputtools.google.com/request?text={}&itc=te-t-i0-und&num=1"
    headers = {
        "User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko)"
    }
    for word in words:
        if word.isdigit() or word in [",", ".", "!", "?", "-", ":", ";"]:
            converted_words.append(word)
            continue
        try:
            req_url = url.format(urllib.parse.quote(word))
            res = requests.get(req_url, headers=headers, timeout=5)
            data = res.json()
            if data[0] == "SUCCESS" and len(data[1][0][1]) > 0:
                converted_words.append(data[1][0][1][0])
            else:
                converted_words.append(word)
        except Exception:
            converted_words.append(word)
    return " ".join(converted_words)

# High-Speed AI Engine
def generate_ai_response(keys: list, contents_list: list, system_instruction: str, temperature=0.3):
    last_error = None
    active_models = ["gemini-2.5-flash", "gemini-2.0-flash", "gemini-3.5-flash-lite", "gemini-3.5-flash"]
    
    for key in keys:
        if not key or not key.strip():
            continue
        try:
            client = genai.Client(api_key=key.strip())
            for model_name in active_models:
                try:
                    res = client.models.generate_content(
                        model=model_name,
                        contents=contents_list,
                        config=types.GenerateContentConfig(
                            system_instruction=system_instruction,
                            temperature=temperature,
                        ),
                    )
                    if res and res.text:
                        return res.text
                except Exception as model_err:
                    last_error = model_err
                    continue
        except Exception as client_err:
            last_error = client_err
            continue
            
    if last_error:
        raise last_error
    raise RuntimeError("API key unavailable or quota exceeded.")

# DOCX Generator
def create_docx_press_note(text: str, date_str: str, location_str: str, is_short: bool) -> io.BytesIO:
    doc = Document()
    p_header = doc.add_paragraph()
    run_name = p_header.add_run("TATA MADHUSUDHAN\n")
    run_name.font.name = "Arial"
    run_name.font.size = Pt(16)
    run_name.bold = True
    run_name.font.color.rgb = RGBColor(88, 28, 135)
    
    run_desig = p_header.add_run("M.L.C\nKhammam, Telangana\n")
    run_desig.font.name = "Arial"
    run_desig.font.size = Pt(11)
    run_desig.bold = True
    
    doc.add_paragraph("―" * 55)
    
    mode_label = "సంక్షిప్త ప్రకటన (Short Press Note)" if is_short else "సమగ్ర అధికారిక ప్రకటన (Full Press Release)"
    p_meta = doc.add_paragraph(f"స్థలం: {location_str} | తేదీ: {date_str} | {mode_label}\n")
    p_meta.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    
    for line in text.split("\n"):
        if line.strip():
            p = doc.add_paragraph(line.strip())
            p.paragraph_format.line_spacing = 1.3
            p.paragraph_format.space_after = Pt(6)
            
    bio = io.BytesIO()
    doc.save(bio)
    bio.seek(0)
    return bio

# Printable HTML Template
def get_printable_letterhead_html(content: str, date_str: str, location_str: str, lh_base64: str) -> str:
    if lh_base64:
        header_html = f"""
        <div style="width: 100%; max-height: 180px; overflow: hidden; border-bottom: 2px solid #581c87; margin-bottom: 20px;">
            <img src="data:image/png;base64,{lh_banner_base64}" style="width: 100%; object-fit: cover; object-position: top center;" />
        </div>
        """
    else:
        header_html = """
        <div style="border-bottom: 2px solid #581c87; padding-bottom: 12px; margin-bottom: 20px;">
            <table style="width: 100%;">
                <tr>
                    <td style="width: 40%; vertical-align: top;">
                        <h2 style="color: #581c87; margin: 0; font-size: 22px; font-weight: 800;">TATA MADHUSUDHAN</h2>
                        <div style="font-size: 15px; font-weight: bold;">M.L.C</div>
                        <div style="font-size: 13px; color: #555;">Khammam, Telangana</div>
                    </td>
                    <td style="width: 20%; text-align: center; vertical-align: middle;">
                        <div style="font-size: 40px;">🏛️</div>
                    </td>
                    <td style="width: 40%; text-align: right; vertical-align: top; font-size: 11px; color: #444; line-height: 1.4;">
                        Quarter No. 1104, 11th Floor,<br>
                        M.S. Block-III, Old MLA Quarters,<br>
                        Hyderguda, Hyderabad - 500029<br>
                        e-mail: tatamadhu@gmail.com
                    </td>
                </tr>
            </table>
        </div>
        """

    clean_content = content.replace("<", "&lt;").replace(">", "&gt;").replace("\n", "<br>")
    return f"""<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>MLC Tata Madhusudhan Press Note</title>
<style>
    @import url('https://fonts.googleapis.com/css2?family=Anek+Telugu:wght@400;600;700;800&display=swap');
    body {{
        font-family: 'Anek Telugu', sans-serif;
        background-color: #ffffff;
        margin: 0;
        padding: 20px;
        color: #111;
    }}
    .container {{
        max-width: 800px;
        margin: 0 auto;
        padding: 10px 20px;
        box-sizing: border-box;
    }}
    .meta {{
        text-align: right;
        font-weight: 600;
        font-size: 15px;
        margin-bottom: 18px;
        color: #374151;
    }}
    .content {{
        font-size: 17px;
        line-height: 1.95;
    }}
    @media print {{
        body {{ padding: 0; }}
        .container {{ border: none; padding: 0; }}
        .no-print {{ display: none; }}
    }}
</style>
</head>
<body>
<div class="no-print" style="text-align: center; margin-bottom: 25px;">
    <button onclick="window.print()" style="padding: 12px 28px; font-size: 16px; font-weight: bold; background-color: #581c87; color: white; border: none; border-radius: 6px; cursor: pointer;">
        🖨️ నేరుగా ప్రింట్ / PDF తీయండి (Print or Save as PDF)
    </button>
</div>
<div class="container">
    {header_html}
    <div class="meta">స్థలం: {location_str} &nbsp;|&nbsp; తేదీ: {date_str}</div>
    <div class="content">
        {clean_content}
    </div>
    <div style="border-top: 1px dashed #581c87; margin-top: 35px; padding-top: 12px; text-align: right; font-weight: bold; color: #444;">
        విడుదల: ఎమ్మెల్సీ తాతా మధుసూదన్ గారి కార్యాలయం
    </div>
</div>
</body>
</html>"""

# ----------------- SIDEBAR -----------------
with st.sidebar:
    st.markdown("### సెట్టింగ్స్ (Settings)")
    default_key = ""
    backup_key = ""
    try:
        if hasattr(st, "secrets"):
            default_key = st.secrets.get("GEMINI_API_KEY", "")
            backup_key = st.secrets.get("BACKUP_API_KEY", "")
    except Exception:
        pass

    api_key_input = st.text_input(
        "Gemini API Key",
        value=default_key,
        type="password",
        help="Google AI Studio API Key ఇక్కడ నమోదు చేయండి."
    )
    
    st.markdown("---")
    dev_img_html = f'<img src="data:image/jpeg;base64,{dev_image_base64}" class="dev-img">' if dev_image_base64 else '<span style="font-size:26px;">👨‍💻</span>'
    st.markdown(f"""
    <div class="dev-badge">
        {dev_img_html}
        <div>
            <div class="dev-text">Designed & Developed by<br>Sumanth Muthamala</div>
            <div class="dev-sub">Revenue Inspector & PA to MLC Khammam</div>
        </div>
    </div>
    """, unsafe_allow_html=True)

active_keys = [api_key_input]
if backup_key and backup_key != api_key_input:
    active_keys.append(backup_key)

# ----------------- MAIN UI -----------------
st.markdown(f"""
<div class="header-card">
    <h1>🎙️ ఎమ్మెల్సీ తాతా మధుసూదన్ — AI ప్రెస్ నోట్ & సందేశాల కన్సోల్</h1>
    <p>శాసనమండలి మీడియా సమన్వయ విభాగం & పార్టీ కమ్యూనికేషన్స్ | BRS Party</p>
</div>
""", unsafe_allow_html=True)

if not api_key_input:
    st.warning("ముందుగా సైడ్‌బార్‌లో మీ Gemini API Keyని నమోదు చేయండి.")
    st.stop()

# Main Console Configuration Bar
with st.container():
    col_c1, col_c2, col_c3 = st.columns([1, 1.4, 1.6])
    
    with col_c1:
        st.markdown("##### 📅 తేదీ (Date):")
        selected_date = st.date_input("తేదీని ఎంచుకోండి", value=datetime.date.today(), label_visibility="collapsed")
        formatted_date = selected_date.strftime("%d-%m-%Y")

    with col_c2:
        st.markdown("##### 📍 స్థలం / వేదిక (Location/Venue):")
        if "final_loc_key" not in st.session_state:
            st.session_state["final_loc_key"] = "ఖమ్మం"
            
        loc_eng = st.text_input(
            "ఇంగ్లీష్‌లో టైప్ చేయండి:", 
            placeholder="e.g. Khammam, Telangana Bhavan...", 
            key="loc_eng_input",
            label_visibility="collapsed"
        )
        if st.button("🔄 తెలుగులోకి మార్చండి", use_container_width=True):
            if loc_eng.strip():
                converted_loc = google_transliterate_telugu(loc_eng)
                st.session_state["final_loc_key"] = converted_loc
                st.rerun()

        final_location = st.text_input(
            "ఫైనల్ స్థలం / వేదిక:", 
            key="final_loc_key",
            label_visibility="collapsed"
        )

    with col_c3:
        st.markdown("##### 🏷️ సందర్భం / విభాగం (Context / Topic Scope):")
        topic_scopes = [
            "పార్టీ ముఖ్య సమావేశం (Party Key Leaders Meeting)",
            "నియోజకవర్గ విస్తృతస్థాయి సమావేశం (Constituency Broad Meeting)",
            "ఆందోళనలు, ధర్నాలు & నిరసనలు (Protests & Agitations)",
            "ప్రజా సమస్యలు & వినతులు (Public Grievances & Demands)",
            "రైతాంగ & వ్యవసాయ సమస్యలు (Farmers & Agriculture Issues)",
            "శుభాకాంక్షలు & సంతాపాలు (Greetings & Condolences)",
            "ప్రభుత్వ విధానాలపై ప్రెస్ మీట్ (Press Meet on Govt Policies)",
            "సేవా కార్యక్రమాలు & సంక్షేమం (Social Welfare & Charity)",
        ]
        selected_scope = st.selectbox("విభాగం ఎంచుకోండి:", topic_scopes, label_visibility="collapsed")

st.divider()

# Communication Category Selector
col_cat1, col_cat2 = st.columns([1.3, 1])

with col_cat1:
    comm_mode = st.radio(
        "🎯 మీరు రూపొందించాలనుకుంటున్న కమ్యూనికేషన్ రకం:",
        ["📢 సమావేశ ఆహ్వానాలు & సందేశాలు (Party Leaders, Cadre & Press)", "📰 అధికారిక పత్రికా ప్రకటన (Press Note)"],
        index=0,
        horizontal=True
    )

with col_cat2:
    if "సందేశాలు" in comm_mode:
        msg_meeting_time = st.text_input("⏰ సమావేశ సమయం (Time):", value="ఉదయం 11:00 గంటలకు", placeholder="e.g., ఉదయం 10:30 గంటలకు")
    else:
        press_length = st.radio(
            "📝 ప్రెస్ నోట్ రకం:",
            ["📰 సమగ్ర ప్రకటన (Long)", "⚡ సంక్షిప్త ప్రకటన (Short)"],
            index=0,
            horizontal=True
        )

# Media & Prompt Input Selector
input_mode = st.radio(
    "📥 ఇన్‌పుట్ విధానం (Input Mode):",
    ["✍️ సిట్యుయేషన్ నోట్స్ / ప్రాంప్ట్ (Situational Text Prompt)", "📁 ఆడియో / వీడియో ఫైల్ (File)", "🎤 లైవ్ రికార్డింగ్ (Mic)"],
    index=0,
    horizontal=True
)

st.write("")

if "audio_file_payload" not in st.session_state:
    st.session_state["audio_file_payload"] = None

input_parts = []

if input_mode == "✍️ సిట్యుయేషన్ నోట్స్ / ప్రాంప్ట్ (Situational Text Prompt)":
    st.markdown("##### 🤖 సిట్యుయేషన్ ప్రాంప్ట్ / వివరాలు (Enter Situation Prompt in English, Tanglish, or Telugu):")
    notes_text = st.text_area(
        "ఇక్కడ మీ ప్రాంప్ట్ లేదా నోట్స్ నమోదు చేయండి (English / Telugu / Tanglish):",
        height=135,
        value=st.session_state.get("final_notes_area", ""),
        key="notes_text_area",
        placeholder="ఉదాహరణ: రేపు ఉదయం 10 గంటలకు తెలంగాణ భవన్‌లో రైతు రుణమాఫీపై నియోజకవర్గ ముఖ్య నేతల అత్యవసర సమావేశం, అందరూ హాజరు కావాలి..."
    )
    if notes_text.strip():
        input_parts.append(types.Part.from_text(text=notes_text))

elif input_mode == "📁 ఆడియో / వీడియో ఫైల్ (File)":
    st.markdown("##### 📁 మొబైల్ లేదా ల్యాప్‌టాప్ రికార్డింగ్ ఫైల్ ఎంచుకోండి:")
    col_f1, col_f2 = st.columns([3, 1])
    with col_f1:
        uploaded_file = st.file_uploader(
            "ఆడియో లేదా వీడియో ఫైల్ ఎంచుకోండి (.m4a, .mp3, .wav, .mp4, etc.):",
            type=None,
            key="pinned_mobile_audio_uploader"
        )
    with col_f2:
        if st.session_state.get("audio_file_payload"):
            if st.button("🗑️ ఫైల్ తొలగించు (Clear File)", use_container_width=True):
                st.session_state["audio_file_payload"] = None
                st.rerun()

    if uploaded_file is not None:
        file_bytes = uploaded_file.read()
        f_name = uploaded_file.name
        fn_low = f_name.lower()

        if fn_low.endswith((".m4a", ".aac")) or "m4a" in (uploaded_file.type or "").lower():
            clean_mime = "audio/mp4"
        elif fn_low.endswith(".mp3"):
            clean_mime = "audio/mp3"
        elif fn_low.endswith(".wav"):
            clean_mime = "audio/wav"
        elif fn_low.endswith((".ogg", ".opus")):
            clean_mime = "audio/ogg"
        elif fn_low.endswith((".mp4", ".m4v")):
            clean_mime = "video/mp4"
        elif fn_low.endswith(".mov"):
            clean_mime = "video/quicktime"
        else:
            clean_mime = "audio/mp4"

        st.session_state["audio_file_payload"] = (f_name, file_bytes, clean_mime)

    if st.session_state.get("audio_file_payload"):
        fn, fb, fm = st.session_state["audio_file_payload"]
        st.success(f"✅ ఫైల్ సిద్ధంగా ఉంది: **{fn}** ({len(fb)/(1024*1024):.2f} MB)")

elif input_mode == "🎤 లైవ్ రికార్డింగ్ (Mic)":
    st.markdown("##### మైక్ ద్వారా మాట్లాడి రికార్డ్ చేయండి:")
    live_audio = st.audio_input("వాయిస్ రికార్డ్ చేయండి")
    if live_audio:
        st.success("✅ ఆడియో రికార్డ్ అయ్యింది!")
        audio_bytes = live_audio.read()
        input_parts.append(types.Part.from_bytes(data=audio_bytes, mime_type="audio/wav"))

st.divider()

# Generation Action Button
btn_label = "⚡ సందేశాలను రూపొందించండి (Generate 3 Dedicated Messages)" if "సందేశాలు" in comm_mode else "🚀 పత్రికా ప్రకటనను రూపొందించండి (Generate Press Note)"

if st.button(btn_label, type="primary", use_container_width=True):
    if input_mode == "📁 ఆడియో / వీడియో ఫైల్ (File)" and st.session_state.get("audio_file_payload"):
        fn, fb, fm = st.session_state["audio_file_payload"]
        if len(fb) > 20 * 1024 * 1024:
            file_ext = os.path.splitext(fn)[1] or (".mp4" if "video" in fm else ".m4a")
            with tempfile.NamedTemporaryFile(delete=False, suffix=file_ext) as tmp:
                tmp.write(fb)
                tmp_path = tmp.name
            with st.spinner(f"పెద్ద మీడియా ఫైల్ ({fn}) ప్రాసెస్ అవుతోంది..."):
                try:
                    c = genai.Client(api_key=active_keys[0])
                    up_ref = c.files.upload(file=tmp_path)
                    if "video" in fm:
                        while up_ref.state.name == "PROCESSING":
                            time.sleep(2)
                            up_ref = c.files.get(name=up_ref.name)
                    input_parts.append(up_ref)
                finally:
                    if os.path.exists(tmp_path):
                        os.remove(tmp_path)
        else:
            input_parts.append(types.Part.from_bytes(data=fb, mime_type=fm))

    if not input_parts:
        st.error("⚠️ దయచేసి వివరాలు లేదా ప్రాంప్ట్ నమోదు చేయండి.")
    else:
        with st.spinner("⚡ AI రూపొందిస్తోంది... దయచేసి వేచి ఉండండి..."):
            try:
                if "సందేశాలు" in comm_mode:
                    prompt_instruction = (
                        f"\nసందర్భం / విభాగం: {selected_scope}\n"
                        f"వేదిక / స్థలం: {final_location}\n"
                        f"తేదీ: {formatted_date}\n"
                        f"సమయం: {msg_meeting_time}\n"
                        "నాయకుల కొరకు, పార్టీ శ్రేణుల కొరకు, మరియు ప్రెస్ ప్రతినిధుల కొరకు "
                        "ఎమ్మెల్సీ తాతా మధుసూదన్ గారి పేరున 3 ప్రత్యేక సందేశాలను పైన తెలిపిన ఫార్మాట్‌లో రూపొందించండి.\n"
                    )
                    chosen_sys = SYSTEM_INSTRUCTION_MESSAGES
                else:
                    is_short = "Short" in press_length
                    chosen_sys = SYSTEM_INSTRUCTION_SHORT if is_short else SYSTEM_INSTRUCTION_LONG
                    prompt_instruction = (
                        f"\nప్రకటన విభాగం: {selected_scope}\n"
                        f"స్థలం: {final_location}\n"
                        f"తేదీ: {formatted_date}\n"
                        f"ఫార్మాట్: {'సంక్షిప్త ప్రకటన' if is_short else 'సమగ్ర అధికారిక ప్రకటన'}\n"
                        "ఎమ్మెల్సీ తాతా మధుసూదన్ గారి అధికారిక పత్రికా ప్రకటనను రూపొందించండి.\n"
                    )

                parts_with_prompt = input_parts + [prompt_instruction]
                ai_output = generate_ai_response(active_keys, parts_with_prompt, chosen_sys)
                
                if "సందేశాలు" in comm_mode:
                    st.session_state["messages_output"] = ai_output
                    st.session_state["is_message_mode"] = True
                else:
                    st.session_state["draft_note"] = ai_output
                    st.session_state["final_note"] = ai_output
                    st.session_state["is_short_selected"] = "Short" in press_length
                    st.session_state["is_message_mode"] = False
                    st.session_state["is_finalized"] = False

            except Exception as e:
                st.error(f"సర్వర్ లోపం: {str(e)}")

# =========================================================================
# OUTPUT VIEW: DEDICATED MEETING & SITUATIONAL MESSAGES (3 SLOTS)
# =========================================================================
if st.session_state.get("is_message_mode", False) and "messages_output" in st.session_state:
    st.divider()
    st.subheader("📢 ఎమ్మెల్సీ తాతా మధు గారి అధికారిక సమావేశ సందేశాలు (3 Dedicated Slots)")
    
    raw_msgs = st.session_state["messages_output"]
    leaders_text = ""
    cadre_text = ""
    press_text = ""
    
    if "===LEADERS_MSG===" in raw_msgs:
        parts_1 = raw_msgs.split("===LEADERS_MSG===")[1]
        if "===CADRE_MSG===" in parts_1:
            leaders_text = parts_1.split("===CADRE_MSG===")[0].strip()
            parts_2 = parts_1.split("===CADRE_MSG===")[1]
            if "===PRESS_INVITE===" in parts_2:
                cadre_text = parts_2.split("===PRESS_INVITE===")[0].strip()
                press_text = parts_2.split("===PRESS_INVITE===")[1].strip()
            else:
                cadre_text = parts_2.strip()
        else:
            leaders_text = parts_1.strip()
    else:
        leaders_text = raw_msgs
        cadre_text = raw_msgs
        press_text = raw_msgs

    tab_m1, tab_m2, tab_m3 = st.tabs([
        "👥 ముఖ్య నాయకులకు (BRS Leaders)", 
        "✊ పార్టీ శ్రేణులు & కార్యకర్తలకు (Cadre & Followers)", 
        "🎙️ మీడియా & పత్రికా ప్రతినిధులకు (Press Slot)"
    ])

    with tab_m1:
        st.markdown("##### 📌 ముఖ్య నాయకులు, ప్రజాప్రతినిధులకు సందేశం:")
        edited_leader = st.text_area("సందేశం పరిశీలించండి / సవరించండి:", value=leaders_text, height=200, key="edit_leader_msg")
        render_copy_button(edited_leader, "📋 నాయకుల సందేశాన్ని కాపీ చేయండి (Copy Text)", "btn_copy_leader")
        
        st.markdown("**సోషల్ మీడియా / మెసేజింగ్ షేరింగ్:**")
        col_s1, col_s2, col_s3 = st.columns(3)
        with col_s1:
            wa_lead = f"https://api.whatsapp.com/send?text={urllib.parse.quote(edited_leader)}"
            st.link_button("📲 WhatsApp లో పంపండి", wa_lead, use_container_width=True)
        with col_s2:
            tg_lead = f"https://t.me/share/url?url=&text={urllib.parse.quote(edited_leader)}"
            st.link_button("✈️ Telegram లో పంపండి", tg_lead, use_container_width=True)
        with col_s3:
            tw_lead = f"https://twitter.com/intent/tweet?text={urllib.parse.quote(edited_leader[:250])}"
            st.link_button("🐦 X (Twitter) లో పోస్ట్ చేయండి", tw_lead, use_container_width=True)

    with tab_m2:
        st.markdown("##### 📌 పార్టీ శ్రేణులు, కార్యకర్తలు & యూత్ విభాగానికి పిలుపు:")
        edited_cadre = st.text_area("సందేశం పరిశీలించండి / సవరించండి:", value=cadre_text, height=200, key="edit_cadre_msg")
        render_copy_button(edited_cadre, "📋 కార్యకర్తల సందేశాన్ని కాపీ చేయండి (Copy Text)", "btn_copy_cadre")
        
        st.markdown("**సోషల్ మీడియా / మెసేజింగ్ షేరింగ్:**")
        col_c1, col_c2, col_c3 = st.columns(3)
        with col_c1:
            wa_cad = f"https://api.whatsapp.com/send?text={urllib.parse.quote(edited_cadre)}"
            st.link_button("📲 WhatsApp లో పంపండి", wa_cad, use_container_width=True)
        with col_c2:
            tg_cad = f"https://t.me/share/url?url=&text={urllib.parse.quote(edited_cadre)}"
            st.link_button("✈️ Telegram లో పంపండి", tg_cad, use_container_width=True)
        with col_c3:
            fb_cad = f"https://www.facebook.com/sharer/sharer.php?u=&quote={urllib.parse.quote(edited_cadre)}"
            st.link_button("🔷 Facebook లో పోస్ట్ చేయండి", fb_cad, use_container_width=True)

    with tab_m3:
        st.markdown("##### 📌 ప్రింట్, ఎలక్ట్రానిక్ & డిజిటల్ మీడియా ప్రతినిధులకు ఆహ్వానం / నోట్:")
        edited_press = st.text_area("మీడియా నోట్ పరిశీలించండి / సవరించండి:", value=press_text, height=200, key="edit_press_msg")
        
        col_pr_act1, col_pr_act2 = st.columns(2)
        with col_pr_act1:
            render_copy_button(edited_press, "📋 మీడియా సందేశాన్ని కాపీ చేయండి (Copy Text)", "btn_copy_press")
        with col_pr_act2:
            if st.button("📄 లెటర్‌హెడ్‌పై ప్రెస్ నోట్‌గా మార్చండి (View on Letterhead)", use_container_width=True):
                st.session_state["draft_note"] = edited_press
                st.session_state["final_note"] = edited_press
                st.session_state["is_short_selected"] = True
                st.session_state["is_message_mode"] = False
                st.session_state["is_finalized"] = True
                st.rerun()

        st.markdown("**సోషల్ మీడియా / మెసేజింగ్ షేరింగ్:**")
        col_p1, col_p2, col_p3 = st.columns(3)
        with col_p1:
            wa_prs = f"https://api.whatsapp.com/send?text={urllib.parse.quote(edited_press)}"
            st.link_button("📲 WhatsApp లో పంపండి", wa_prs, use_container_width=True)
        with col_p2:
            tw_prs = f"https://twitter.com/intent/tweet?text={urllib.parse.quote(edited_press[:250])}"
            st.link_button("🐦 X (Twitter) లో షేర్ చేయండి", tw_prs, use_container_width=True)
        with col_p3:
            tg_prs = f"https://t.me/share/url?url=&text={urllib.parse.quote(edited_press)}"
            st.link_button("✈️ Telegram లో షేర్ చేయండి", tg_prs, use_container_width=True)

# =========================================================================
# OUTPUT VIEW: OFFICIAL PRESS NOTE (ON OFFICIAL LETTERHEAD CANVAS)
# =========================================================================
elif not st.session_state.get("is_message_mode", False) and "draft_note" in st.session_state:
    st.subheader("✏️ ఎడిట్ & AI సలహాలు (Edit & AI Remarks)")
    
    edited_note = st.text_area(
        "ముసాయిదాను ఇక్కడ పరిశీలించి నేరుగా సవరించవచ్చు:",
        value=st.session_state.get("draft_note", ""),
        height=260,
        key="editor_area"
    )
    
    st.markdown("##### 🤖 జెమినీ AI సలహా / మోడ్ మార్పు:")
    with st.expander("💡 ప్రెస్ నోట్ మార్పులు (రైతుల ఆవేదన పెంచడం, వివరాలు చేర్చడం, స్పష్టత ఇవ్వడం)", expanded=False):
        ai_remark = st.text_input(
            "మీ సూచనను ఇక్కడ రాయండి (English or Telugu):",
            placeholder="e.g., 'సంతాప సందేశాన్ని మరింత భావోద్వేగంగా మార్చండి'..."
        )
        if st.button("⚡ సూచన ఆధారంగా ప్రెస్ నోట్ తిరిగి రూపొందించండి (Re-generate with AI)", use_container_width=True):
            if ai_remark.strip():
                with st.spinner("సరిచేస్తోంది..."):
                    try:
                        is_short = st.session_state.get("is_short_selected", False)
                        chosen_sys = SYSTEM_INSTRUCTION_SHORT if is_short else SYSTEM_INSTRUCTION_LONG
                        refine_prompt = f"""
                        CURRENT PRESS NOTE DRAFT:
                        {edited_note}
                        
                        USER INSTRUCTION:
                        {ai_remark}
                        
                        TASK:
                        Rewrite and refine strictly following user instructions while maintaining MLC Tata Madhusudhan Garu's persona.
                        """
                        updated_note = generate_ai_response(active_keys, [refine_prompt], chosen_sys)
                        st.session_state["draft_note"] = updated_note
                        st.session_state["final_note"] = updated_note
                        st.success("✅ విజయవంతంగా మార్చబడింది!")
                        st.rerun()
                    except Exception as err:
                        st.error(f"రీ-జనరేట్ ఎర్రర్: {str(err)}")

    st.write("")
    if st.button("✅ పూర్తయింది - లెటర్‌హెడ్ & ఎగుమతి వీక్షించండి (Finalize)", type="primary", use_container_width=True):
        st.session_state["final_note"] = edited_note
        st.session_state["is_finalized"] = True
        st.success("ప్రెస్ నోట్ ఖరారైంది! అధికారిక లెటర్‌హెడ్ వీక్షణ సిద్ధంగా ఉంది.")

if st.session_state.get("is_finalized", False) and not st.session_state.get("is_message_mode", False):
    final_content = st.session_state.get("final_note", "")
    is_short = st.session_state.get("is_short_selected", False)
    
    lines = [line.strip() for line in final_content.split("\n") if line.strip()]
    headline = lines[0] if lines else "ఎమ్మెల్సీ తాతా మధుసూదన్ గారి ప్రకటన"
    for l in lines:
        if "హెడ్" in l or "శీర్షిక" in l or ":" in l:
            headline = l.split(":")[-1].strip().replace("*", "")
            break

    st.divider()
    st.markdown("### 📋 ప్రెస్ నోట్ కాపీ & ఎగుమతి (Copy & Export)")
    
    render_copy_button(final_content, "📋 పూర్తి ప్రెస్ నోట్ కాపీ చేయండి (Copy Full Press Note)", "btn_copy_final_press")
    
    col_exp1, col_exp2 = st.columns([1, 1])
    with col_exp1:
        docx_data = create_docx_press_note(final_content, formatted_date, final_location, is_short)
        st.download_button(
            label="📄 Word File (.DOCX) డౌన్‌లోడ్ చేయండి",
            data=docx_data,
            file_name=f"Tata_Madhu_{formatted_date}.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            use_container_width=True
        )
    with col_exp2:
        html_page = get_printable_letterhead_html(final_content, formatted_date, final_location, lh_banner_base64)
        st.download_button(
            label="🖨️ PDF / Print File (.HTML) డౌన్‌లోడ్ చేయండి",
            data=html_page,
            file_name=f"Tata_Madhu_Letterhead_{formatted_date}.html",
            mime="text/html",
            use_container_width=True
        )

    st.write("")
    st.subheader("📄 అధికారిక లెటర్‌హెడ్ వీక్షణ (Official Letterhead View)")
    
    if lh_banner_base64:
        banner_img_html = f'<img src="data:image/png;base64,{lh_banner_base64}" class="letterhead-banner-img" alt="Official Letterhead">'
    else:
        banner_img_html = """
        <div style="border-bottom: 2px solid #581c87; padding-bottom: 12px; margin-bottom: 20px;">
            <table style="width: 100%;">
                <tr>
                    <td style="width: 40%; vertical-align: top;">
                        <h2 style="color: #581c87; margin: 0; font-size: 22px; font-weight: 800;">TATA MADHUSUDHAN</h2>
                        <div style="font-size: 15px; font-weight: bold;">M.L.C</div>
                        <div style="font-size: 13px; color: #555;">Khammam, Telangana</div>
                    </td>
                    <td style="width: 20%; text-align: center; vertical-align: middle;">
                        <div style="font-size: 40px;">🏛️</div>
                    </td>
                    <td style="width: 40%; text-align: right; vertical-align: top; font-size: 11px; color: #444; line-height: 1.4;">
                        Quarter No. 1104, 11th Floor,<br>
                        M.S. Block-III, Old MLA Quarters,<br>
                        Hyderguda, Hyderabad - 500029<br>
                        e-mail: tatamadhu@gmail.com
                    </td>
                </tr>
            </table>
        </div>
        """

    formatted_body = (
        final_content.replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace("**", "")
        .replace("\n", "<br>")
    )

    canvas_html = f"""<div class="letterhead-container">
{banner_img_html}
<div style="text-align: right; font-weight: 600; font-size: 15px; margin-bottom: 20px; color: #374151;">
స్థలం: {final_location} &nbsp;|&nbsp; తేదీ: {formatted_date}
</div>
<div style="line-height: 1.95; font-size: 18px;">
{formatted_body}
</div>
<div style="border-top: 1.5px dashed #581c87; margin-top: 30px; padding-top: 14px; text-align: right; font-size: 15px; font-weight: 700; color: #374151;">
విడుదల: ఎమ్మెల్సీ తాతా మధుసూదన్ గారి కార్యాలయం
</div>
</div>"""

    st.markdown(canvas_html, unsafe_allow_html=True)
    
    st.write("")
    st.subheader("🌐 సోషల్ మీడియా పోస్టులు (Ready to Share & Post)")
    summary_body = "\n".join(lines[1:5]) if len(lines) > 1 else final_content

    whatsapp_text = f"*{headline}*\n\n{final_content}\n\n_విడుదల: ఎమ్మెల్సీ తాతా మధుసూదన్ గారి కార్యాలయం_"
    twitter_text = f"🚨 {headline[:180]}\n\n- ఎమ్మెల్సీ తాతా మధుసూదన్\n\n#TataMadhu #BRSParty #Telangana #Khammam"
    facebook_text = f"📌 {headline}\n\n{final_content}\n\n#TataMadhusudhan #TataMadhu #BRS #Khammam #TelanganaPolitics"
    telegram_text = f"📢 *{headline}*\n\n{final_content}"

    st1, st2, st3, st4 = st.tabs(["🟢 WhatsApp", "🔵 Twitter (X)", "🔷 Facebook", "✈️ Telegram"])
    with st1:
        st.text_area("WhatsApp Text:", value=whatsapp_text, height=160)
        render_copy_button(whatsapp_text, "📋 WhatsApp టెక్స్ట్ కాపీ చేయండి", "copy_wa")
        wa_url = f"https://api.whatsapp.com/send?text={urllib.parse.quote(whatsapp_text[:1400])}"
        st.link_button("📲 వాట్సాప్‌లో షేర్ చేయండి", wa_url, use_container_width=True)
    with st2:
        st.text_area("Twitter (X) Post:", value=twitter_text, height=130)
        render_copy_button(twitter_text, "📋 X (Twitter) పోస్ట్ కాపీ చేయండి", "copy_tw")
        x_url = f"https://twitter.com/intent/tweet?text={urllib.parse.quote(twitter_text)}"
        st.link_button("🐦 X (Twitter) లో పోస్ట్ చేయండి", x_url, use_container_width=True)
    with st3:
        st.text_area("Facebook Post:", value=facebook_text, height=160)
        render_copy_button(facebook_text, "📋 Facebook పోస్ట్ కాపీ చేయండి", "copy_fb")
        fb_url = f"https://www.facebook.com/sharer/sharer.php?u=&quote={urllib.parse.quote(facebook_text[:1000])}"
        st.link_button("🔷 Facebook లో పోస్ట్ చేయండి", fb_url, use_container_width=True)
    with st4:
        st.text_area("Telegram Post:", value=telegram_text, height=160)
        render_copy_button(telegram_text, "📋 Telegram టెక్స్ట్ కాపీ చేయండి", "copy_tg")
        tg_url = f"https://t.me/share/url?url=&text={urllib.parse.quote(telegram_text[:1400])}"
        st.link_button("✈️ Telegram లో షేర్ చేయండి", tg_url, use_container_width=True)

# Footer
st.markdown("---")
dev_footer_img = f'<img src="data:image/jpeg;base64,{dev_image_base64}" style="width: 36px; height: 36px; border-radius: 50%; vertical-align: middle; margin-right: 10px; border: 1.5px solid #7e22ce;">' if dev_image_base64 else '👨‍💻 '
st.markdown(
    f"""
    <div style="text-align: center; color: #374151; font-size: 14px; padding: 18px 0;">
        {dev_footer_img}
        <strong>Designed & Developed by Sumanth Muthamala</strong> &nbsp;|&nbsp; 
        <span style="color: #6b21a8; font-weight: 600;">Revenue Inspector & PA to MLC Khammam</span>
    </div>
    """, 
    unsafe_allow_html=True
)
